# -*- coding: utf-8 -*-
"""Baut die Kalkulationsmappe für das Fresca-/Grapefruit-Limonaden-Projekt."""
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = "/home/user/aktien-cockpit/Fresca-Projekt/02_Kalkulation.xlsx"

FONT = "Arial"
BLUE, BLACK, GREEN, GREY_T = "0000FF", "000000", "008000", "808080"
YELLOW = PatternFill("solid", fgColor="FFFF00")
GREYF = PatternFill("solid", fgColor="F2F2F2")
DARK = PatternFill("solid", fgColor="1F3864")
MID = PatternFill("solid", fgColor="D9E2F3")

EUR3, EUR2, PCT, NUM, NUM1 = '#,##0.000 €', '#,##0.00 €', '0.0%', '#,##0', '#,##0.0'
thin = Side(style="thin", color="BFBFBF")
BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
TOPLINE = Border(top=Side(style="medium", color="1F3864"))

wb = openpyxl.Workbook()


def title(ws, text, width=8):
    ws.merge_cells(start_row=1, start_column=1, end_row=1, end_column=width)
    c = ws.cell(row=1, column=1, value=text)
    c.font = Font(name=FONT, size=14, bold=True, color="FFFFFF")
    c.fill = DARK
    c.alignment = Alignment(vertical="center", indent=1)
    ws.row_dimensions[1].height = 26


def section(ws, row, text, width=8):
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=width)
    c = ws.cell(row=row, column=1, value=text)
    c.font = Font(name=FONT, size=10, bold=True, color="1F3864")
    c.fill = MID
    c.alignment = Alignment(vertical="center", indent=1)


def put(ws, row, col, value, *, fmt=None, color=BLACK, bold=False, inp=False,
        italic=False, wrap=False, size=10, align=None):
    c = ws.cell(row=row, column=col, value=value)
    c.font = Font(name=FONT, size=size, bold=bold, color=color, italic=italic)
    if fmt:
        c.number_format = fmt
    if inp:
        c.fill = YELLOW
        c.border = BOX
    if wrap:
        c.alignment = Alignment(wrap_text=True, vertical="top")
    if align:
        c.alignment = Alignment(horizontal=align)
    return c


def note(ws, row, col, text):
    return put(ws, row, col, text, italic=True, size=9, color=GREY_T, wrap=True)


def header_row(ws, row, labels, start=1):
    for i, lab in enumerate(labels):
        c = ws.cell(row=row, column=start + i, value=lab)
        c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="4472C4")
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 30


def widths(ws, spec):
    for col, w in spec.items():
        ws.column_dimensions[col].width = w


# ================================================================= ANLEITUNG
ws = wb.active
ws.title = "Anleitung"
title(ws, "FRESCA-PROJEKT — KALKULATION EIGENE GRAPEFRUIT-LIMONADE", 6)
widths(ws, {"A": 32, "B": 88})

def line(row, a, b, head=False):
    if head:
        section(ws, row, a, 2)
    else:
        put(ws, row, 1, a, bold=True, color="1F3864")
        put(ws, row, 2, b, wrap=True)
        ws.row_dimensions[row].height = max(15, 13 * (1 + len(b) // 100))

line(3, "WOZU DIESE MAPPE", None, head=True)
put(ws, 4, 2, "Sie rechnet durch, was dein Getränk kostet, was du daran verdienst und ab wann "
              "sich die erste Charge trägt. Alles hängt zusammen: Änderst du eine Annahme, "
              "rechnet sich die ganze Mappe neu.", wrap=True)
ws.row_dimensions[4].height = 30

line(6, "SO ARBEITEST DU DAMIT", None, head=True)
steps = [
    ("Schritt 1", "Öffne das Blatt 'Annahmen'. NUR dort trägst du etwas ein — alle gelben Felder."),
    ("Schritt 2", "Ersetze die Startwerte durch das, was in den echten Angeboten von Abfüller "
                  "und Aromahaus steht. Die Startwerte sind belastbare Marktschätzungen — aber eben Schätzungen."),
    ("Schritt 3", "Auf 'Rezepturvarianten' stellst du Zucker, Süßstoff und Saftanteil je Variante ein."),
    ("Schritt 4", "Alle anderen Blätter rechnen automatisch. Dort bitte nichts überschreiben."),
]
for i, (a, b) in enumerate(steps):
    line(7 + i, a, b)

line(12, "FARBLEGENDE", None, head=True)
put(ws, 13, 1, "Gelb hinterlegt", bold=True, color="1F3864", inp=True)
put(ws, 13, 2, "Dein Eingabefeld. Nur diese Zellen ändern.")
put(ws, 14, 1, "Blaue Schrift", bold=True, color=BLUE)
put(ws, 14, 2, "Fest eingetragener Wert — eine Annahme oder eine Zahl aus einem Angebot.")
put(ws, 15, 1, "Schwarze Schrift", bold=True, color=BLACK)
put(ws, 15, 2, "Berechnet. Nicht überschreiben.")
put(ws, 16, 1, "Grüne Schrift", bold=True, color=GREEN)
put(ws, 16, 2, "Verweis auf ein anderes Blatt dieser Mappe.")

line(18, "DIE BLÄTTER", None, head=True)
sheets = [
    ("Annahmen", "Alle Eingaben: Stückkosten, Rohstoffpreise, Einmalkosten, Fixkosten, Verkaufspreise."),
    ("Rezepturvarianten", "A = Vollzucker · B = Fresca-Original (Hybrid) · C = nur Zucker, reduziert · D = wie C, aber 6 % Saft."),
    ("Stückkosten", "Was eine Dose kostet — je Variante und je Chargengröße."),
    ("Marge & Preiskette", "Was du je Absatzkanal verdienst (Online, Gastro, Fachhandel, LEH)."),
    ("Break-even", "Ab wie vielen verkauften Dosen die erste Charge im Plus ist."),
    ("Jahresszenarien", "Vom Nebenerwerb bis zum Hauptberuf — was welcher Absatz bedeutet."),
    ("Verkostung", "Auswertung der Blindverkostung. Trag die Papierbögen hier ein."),
]
for i, (a, b) in enumerate(sheets):
    line(19 + i, a, b)

line(27, "WICHTIG", None, head=True)
put(ws, 28, 2, "Die Startwerte stammen aus Marktrecherche, nicht aus Angeboten. Deine erste "
               "Aufgabe ist, sie durch echte Zahlen zu ersetzen. Bis dahin ist jedes Ergebnis "
               "hier eine Indikation — keine Entscheidungsgrundlage.", wrap=True, color="C00000")
ws.row_dimensions[28].height = 32

# ================================================================= ANNAHMEN
ws = wb.create_sheet("Annahmen")
title(ws, "ANNAHMEN — nur die gelben Felder ändern", 6)
widths(ws, {"A": 48, "B": 16, "C": 16, "D": 16, "E": 3, "F": 62})

section(ws, 3, "1. PROJEKT-GRUNDDATEN")
put(ws, 4, 1, "Chargengröße (Dosen)")
put(ws, 4, 2, 20000, fmt=NUM, color=BLUE, inp=True)
put(ws, 5, 1, "Gebindegröße (ml)")
put(ws, 5, 2, 330, fmt=NUM, color=BLUE, inp=True)
put(ws, 6, 1, "Gewählte Rezepturvariante (A / B / C / D)")
put(ws, 6, 2, "B", color=BLUE, inp=True, align="center")
note(ws, 6, 6, "Steuert die Blätter 'Marge & Preiskette', 'Break-even' und 'Jahresszenarien'.")

section(ws, 8, "2. STÜCKKOSTEN NACH CHARGENGRÖSSE (€ je Dose) — aus den Abfüller-Angeboten eintragen")
header_row(ws, 9, ["Kostenposition", None, None, None])
# Die Mengenstufen MÜSSEN Zahlen sein — sonst findet MATCH() auf anderen
# Blättern die richtige Kostenstufe nicht.
for j, v in enumerate((20000, 50000, 100000)):
    c = ws.cell(row=9, column=2 + j, value=v)
    c.font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="4472C4")
    c.alignment = Alignment(horizontal="center", vertical="center")
    c.number_format = '#,##0" Dosen"'
scale = [
    ("Leerdose + Deckel", 0.150, 0.135, 0.120),
    ("Sleeve (Druck + Aufbringen)", 0.050, 0.040, 0.030),
    ("Abfüllung (Lohnkosten)", 0.120, 0.090, 0.070),
    ("Sekundärverpackung (Tray, Karton, Palette)", 0.030, 0.025, 0.020),
    ("Basis-Rohstoffe (Wasser, CO2, Säure, Konservierung)", 0.020, 0.018, 0.016),
    ("Aroma / Konzentrat", 0.025, 0.022, 0.020),
]
for i, (lab, a, b, c) in enumerate(scale):
    rr = 10 + i
    put(ws, rr, 1, lab)
    for j, v in enumerate((a, b, c)):
        put(ws, rr, 2 + j, v, fmt=EUR3, color=BLUE, inp=True)
note(ws, 10, 6, "Genau diese sechs Zeilen beim Abfüller anfragen — dann sind Angebote vergleichbar.")

section(ws, 17, "3. REZEPTURABHÄNGIGE ROHSTOFFPREISE")
put(ws, 18, 1, "Zucker (€ je kg)")
put(ws, 18, 2, 0.70, fmt=EUR2, color=BLUE, inp=True)
note(ws, 18, 6, "EU-Weißzucker im Großhandel. Schwankt stark — aktuellen Preis erfragen.")
put(ws, 19, 1, "Grapefruitsaft aus Konzentrat (€ je Liter Fertigsaft)")
put(ws, 19, 2, 1.60, fmt=EUR2, color=BLUE, inp=True)
put(ws, 20, 1, "Süßstoff-Blend Sucralose + Acesulfam-K (€ je Dose)")
put(ws, 20, 2, 0.010, fmt=EUR3, color=BLUE, inp=True)

section(ws, 22, "4. ZUSCHLÄGE JE DOSE (€)")
put(ws, 23, 1, "EWP Pfand-Systembeitrag (Produzentenbeitrag)")
put(ws, 23, 2, 0.020, fmt=EUR3, color=BLUE, inp=True)
note(ws, 23, 6, "Tarif bei EWP Recycling Pfand Österreich erfragen. NICHT das 25-Cent-Pfand selbst — das ist ein durchlaufender Posten.")
put(ws, 24, 1, "Fracht Abfüller → eigenes Lager")
put(ws, 24, 2, 0.020, fmt=EUR3, color=BLUE, inp=True)

section(ws, 26, "5. EINMALKOSTEN (€) — fallen nur einmal vor der ersten Charge an")
once = [("Rezepturentwicklung Aromahaus", 3500),
        ("Nährwert- & Stabilitätsanalyse (Labor)", 600),
        ("Dosendesign / Grafik", 1500),
        ("Markenanmeldung EUIPO (1 Klasse, Kl. 32)", 850),
        ("Gewerbeanmeldung, LMU-Registrierung, HACCP-Konzept", 650),
        ("EWP- und ARA-Registrierung (Setup)", 350)]
for i, (lab, v) in enumerate(once):
    put(ws, 27 + i, 1, lab)
    put(ws, 27 + i, 2, v, fmt=EUR2, color=BLUE, inp=True)
put(ws, 33, 1, "Summe Einmalkosten", bold=True)
put(ws, 33, 2, "=SUM(B27:B32)", fmt=EUR2, bold=True)
ws["A33"].border = TOPLINE; ws["B33"].border = TOPLINE

section(ws, 35, "6. FIXKOSTEN JAHR 1 (€)")
for i, (lab, v) in enumerate([("Lagermiete", 1200), ("Logistik / Versand", 1500),
                              ("Marketing, Muster, Messen", 2500),
                              ("Buchhaltung, Versicherung, Sonstiges", 800)]):
    put(ws, 36 + i, 1, lab)
    put(ws, 36 + i, 2, v, fmt=EUR2, color=BLUE, inp=True)
put(ws, 40, 1, "Summe Fixkosten Jahr 1", bold=True)
put(ws, 40, 2, "=SUM(B36:B39)", fmt=EUR2, bold=True)
ws["A40"].border = TOPLINE; ws["B40"].border = TOPLINE

section(ws, 42, "7. ABSATZKANÄLE — Nettopreise, ohne 20 % USt und ohne Pfand")
header_row(ws, 43, ["Kanal", "Abgabepreis netto (€/Dose)", "Anteil am Absatz"])
for i, (lab, p, m) in enumerate([("Online-Direktverkauf", 2.20, 0.25),
                                 ("Gastronomie", 1.00, 0.60),
                                 ("Feinkost / Getränkefachhandel", 0.85, 0.15),
                                 ("LEH / Großhandel", 0.72, 0.00)]):
    put(ws, 44 + i, 1, lab)
    put(ws, 44 + i, 2, p, fmt=EUR2, color=BLUE, inp=True)
    put(ws, 44 + i, 3, m, fmt=PCT, color=BLUE, inp=True)
put(ws, 48, 1, "Summe Anteile (muss 100,0 % ergeben)", bold=True)
put(ws, 48, 3, "=SUM(C44:C47)", fmt=PCT, bold=True)
ws["A48"].border = TOPLINE; ws["C48"].border = TOPLINE
put(ws, 48, 6, '=IF(ROUND(C48,4)=1,"OK","ACHTUNG: Anteile ergeben nicht 100 % — alle Ergebnisse verzerrt!")',
    bold=True, color="C00000")

# ========================================================= REZEPTURVARIANTEN
ws = wb.create_sheet("Rezepturvarianten")
title(ws, "REZEPTURVARIANTEN — Zucker, Süßstoff, Saftanteil", 7)
widths(ws, {"A": 44, "B": 16, "C": 16, "D": 16, "E": 16, "F": 3, "G": 54})

header_row(ws, 3, ["Parameter", "A", "B", "C", "D"])
put(ws, 4, 1, "Bezeichnung", bold=True)
for i, n in enumerate(["Vollzucker", "Fresca-Original (Hybrid)",
                       "Nur Zucker, reduziert", "Nur Zucker + 6 % Saft"]):
    put(ws, 4, 2 + i, n, bold=True, wrap=True, size=9)
ws.row_dimensions[4].height = 30

section(ws, 6, "EINGABEN", 5)
put(ws, 7, 1, "Zucker (g je 100 ml)")
for i, v in enumerate([10, 6, 6, 6]):
    put(ws, 7, 2 + i, v, fmt=NUM1, color=BLUE, inp=True)
put(ws, 8, 1, "Süßstoff enthalten (1 = ja, 0 = nein)")
for i, v in enumerate([0, 1, 0, 0]):
    put(ws, 8, 2 + i, v, fmt=NUM, color=BLUE, inp=True)
put(ws, 9, 1, "Grapefruitsaft-Anteil")
for i, v in enumerate([0.01, 0.01, 0.01, 0.06]):
    put(ws, 9, 2 + i, v, fmt=PCT, color=BLUE, inp=True)
note(ws, 7, 7, "Das Costa-Rica-Original liegt bei ca. 5,9 g/100 ml — Coca-Cola bei ca. 10,6 g.")
note(ws, 9, 7, "Ob 'Fruchtlimonade' auf die Dose darf, hängt vom Mindestfruchtgehalt ab — Österr. Lebensmittelbuch Kap. B26, prüfen lassen.")
ws.row_dimensions[9].height = 30

section(ws, 11, "BERECHNETE REZEPTURKOSTEN (€ je Dose)", 5)
for rr, lab in [(12, "Zucker je Dose (g)"), (13, "Zuckerkosten"),
                (14, "Saftmenge je Dose (ml)"), (15, "Saftkosten"), (16, "Süßstoffkosten")]:
    put(ws, rr, 1, lab)
for i in range(4):
    col = get_column_letter(2 + i)
    put(ws, 12, 2 + i, f"={col}7*Annahmen!$B$5/100", fmt=NUM1)
    put(ws, 13, 2 + i, f"={col}12/1000*Annahmen!$B$18", fmt=EUR3)
    put(ws, 14, 2 + i, f"=Annahmen!$B$5*{col}9", fmt=NUM1)
    put(ws, 15, 2 + i, f"={col}14/1000*Annahmen!$B$19", fmt=EUR3)
    put(ws, 16, 2 + i, f"={col}8*Annahmen!$B$20", fmt=EUR3)

put(ws, 17, 1, "Rezepturkosten gesamt (€ je Dose)", bold=True)
for i in range(4):
    col = get_column_letter(2 + i)
    put(ws, 17, 2 + i, f"=SUM({col}13,{col}15,{col}16)", fmt=EUR3, bold=True)
    ws[f"{col}17"].border = TOPLINE
ws["A17"].border = TOPLINE

section(ws, 19, "NÄHRWERTE UND POSITIONIERUNG", 5)
for rr, lab in [(20, "Brennwert (kcal je 100 ml)"), (21, "Brennwert je Dose (kcal)"),
                (22, "Zucker je Dose (g)")]:
    put(ws, rr, 1, lab)
for i in range(4):
    col = get_column_letter(2 + i)
    put(ws, 20, 2 + i, f"={col}7*4", fmt=NUM1)
    put(ws, 21, 2 + i, f"={col}20*Annahmen!$B$5/100", fmt=NUM)
    put(ws, 22, 2 + i, f"={col}12", fmt=NUM1, color=GREEN)
note(ws, 20, 7, "Zucker liefert 4 kcal je Gramm. Süßstoffe liefern keine Kalorien.")
note(ws, 21, 7, "Vergleich: Coca-Cola hat ca. 139 kcal je 330-ml-Dose.")

# ============================================================== STÜCKKOSTEN
ws = wb.create_sheet("Stückkosten")
title(ws, "STÜCKKOSTEN JE DOSE", 7)
widths(ws, {"A": 44, "B": 16, "C": 16, "D": 16, "E": 16, "F": 3, "G": 54})

put(ws, 3, 1, "Verwendete Kostenstufe (Dosen)", bold=True)
put(ws, 3, 2, "=IFERROR(INDEX(Annahmen!$B$9:$D$9,MATCH(Annahmen!$B$4,Annahmen!$B$9:$D$9,1)),Annahmen!$B$9)",
    fmt=NUM, bold=True, color=GREEN)
note(ws, 3, 7, "Richtet sich nach der Chargengröße auf 'Annahmen'. Zwischenwerte nutzen die nächstniedrigere Stufe.")

header_row(ws, 5, ["Kostenposition", "A", "B", "C", "D"])
put(ws, 6, 1, "Bezeichnung", bold=True)
for i in range(4):
    col = get_column_letter(2 + i)
    put(ws, 6, 2 + i, f"=Rezepturvarianten!{col}4", bold=True, wrap=True, size=9, color=GREEN)
ws.row_dimensions[6].height = 30

for rr, lab, src in [(7, "Leerdose + Deckel", 10), (8, "Sleeve (Druck + Aufbringen)", 11),
                     (9, "Abfüllung (Lohnkosten)", 12), (10, "Sekundärverpackung", 13),
                     (11, "Basis-Rohstoffe", 14), (12, "Aroma / Konzentrat", 15)]:
    put(ws, rr, 1, lab)
    for i in range(4):
        put(ws, rr, 2 + i,
            f"=IFERROR(INDEX(Annahmen!$B${src}:$D${src},MATCH(Annahmen!$B$4,Annahmen!$B$9:$D$9,1)),Annahmen!$B${src})",
            fmt=EUR3)

for rr, lab, src in [(13, "Zucker", 13), (14, "Grapefruitsaft", 15), (15, "Süßstoff", 16)]:
    put(ws, rr, 1, lab)
    for i in range(4):
        col = get_column_letter(2 + i)
        put(ws, rr, 2 + i, f"=Rezepturvarianten!{col}{src}", fmt=EUR3, color=GREEN)

put(ws, 16, 1, "Herstellkosten ab Abfüller (Summe)", bold=True)
for i in range(4):
    col = get_column_letter(2 + i)
    put(ws, 16, 2 + i, f"=SUM({col}7:{col}15)", fmt=EUR3, bold=True)
    ws[f"{col}16"].border = TOPLINE
ws["A16"].border = TOPLINE

put(ws, 17, 1, "+ EWP Pfand-Systembeitrag")
put(ws, 18, 1, "+ Fracht zum Lager")
for i in range(4):
    put(ws, 17, 2 + i, "=Annahmen!$B$23", fmt=EUR3, color=GREEN)
    put(ws, 18, 2 + i, "=Annahmen!$B$24", fmt=EUR3, color=GREEN)

put(ws, 19, 1, "VOLLKOSTEN WARE (€ je Dose)", bold=True)
ws["A19"].fill = MID; ws["A19"].border = TOPLINE
for i in range(4):
    col = get_column_letter(2 + i)
    c = put(ws, 19, 2 + i, f"={col}16+{col}17+{col}18", fmt=EUR3, bold=True)
    c.fill = MID; c.border = TOPLINE

put(ws, 21, 1, "Nachrichtlich: kcal je Dose")
put(ws, 22, 1, "Nachrichtlich: Zucker je Dose (g)")
for i in range(4):
    col = get_column_letter(2 + i)
    put(ws, 21, 2 + i, f"=Rezepturvarianten!{col}21", fmt=NUM, color=GREEN)
    put(ws, 22, 2 + i, f"=Rezepturvarianten!{col}22", fmt=NUM1, color=GREEN)

section(ws, 24, "SKALENEFFEKT — Vollkosten Ware (€ je Dose) bei allen drei Chargengrößen", 5)
header_row(ws, 25, ["Chargengröße", "A", "B", "C", "D"])
for rr, sc in [(26, "B"), (27, "C"), (28, "D")]:
    put(ws, rr, 1, f"=Annahmen!${sc}$9", fmt=NUM, color=GREEN)
    for i in range(4):
        col = get_column_letter(2 + i)
        put(ws, rr, 2 + i,
            f"=SUM(Annahmen!${sc}$10:${sc}$15)+Rezepturvarianten!{col}$17+Annahmen!$B$23+Annahmen!$B$24",
            fmt=EUR3)
note(ws, 26, 7, "Zeigt, wie stark die Menge auf die Stückkosten drückt — dein wichtigstes Argument in der Abfüller-Verhandlung.")
ws.row_dimensions[26].height = 30

# ==================================================== MARGE & PREISKETTE
ws = wb.create_sheet("Marge & Preiskette")
title(ws, "MARGE & PREISKETTE", 8)
widths(ws, {"A": 34, "B": 17, "C": 17, "D": 17, "E": 12, "F": 12, "G": 19, "H": 3, "I": 52})

put(ws, 3, 1, "Gewählte Variante", bold=True)
put(ws, 3, 2, "=Annahmen!$B$6", bold=True, color=GREEN, align="center")
put(ws, 4, 1, "Bezeichnung")
put(ws, 4, 2, "=IFERROR(INDEX(Rezepturvarianten!$B$4:$E$4,MATCH(Annahmen!$B$6,'Stückkosten'!$B$5:$E$5,0)),\"?\")",
    color=GREEN)
put(ws, 5, 1, "Vollkosten Ware (€ je Dose)", bold=True)
put(ws, 5, 2, "=IFERROR(INDEX('Stückkosten'!$B$19:$E$19,MATCH(Annahmen!$B$6,'Stückkosten'!$B$5:$E$5,0)),0)",
    fmt=EUR3, bold=True, color=GREEN)

header_row(ws, 7, ["Kanal", "Abgabepreis netto", "Vollkosten Ware", "Rohertrag € / Dose",
                   "Marge %", "Anteil", "Gewichteter Rohertrag"])
for i in range(4):
    rr, src = 8 + i, 44 + i
    put(ws, rr, 1, f"=Annahmen!$A${src}", color=GREEN)
    put(ws, rr, 2, f"=Annahmen!$B${src}", fmt=EUR2, color=GREEN)
    put(ws, rr, 3, "=$B$5", fmt=EUR3, color=GREEN)
    put(ws, rr, 4, f"=B{rr}-C{rr}", fmt=EUR2, bold=True)
    put(ws, rr, 5, f"=IFERROR(D{rr}/B{rr},0)", fmt=PCT)
    put(ws, rr, 6, f"=Annahmen!$C${src}", fmt=PCT, color=GREEN)
    put(ws, rr, 7, f"=D{rr}*F{rr}", fmt=EUR2)

put(ws, 12, 1, "GEWICHTETER DURCHSCHNITT", bold=True)
put(ws, 12, 2, "=SUMPRODUCT($B$8:$B$11,$F$8:$F$11)", fmt=EUR2, bold=True)
put(ws, 12, 3, "=$B$5", fmt=EUR3, bold=True)
put(ws, 12, 4, "=SUM(G8:G11)", fmt=EUR2, bold=True)
put(ws, 12, 5, "=IFERROR(D12/B12,0)", fmt=PCT, bold=True)
put(ws, 12, 6, "=SUM(F8:F11)", fmt=PCT, bold=True)
put(ws, 12, 7, "=SUM(G8:G11)", fmt=EUR2, bold=True)
for col in "ABCDEFG":
    ws[f"{col}12"].border = TOPLINE
    ws[f"{col}12"].fill = MID

put(ws, 14, 1, "Ø Abgabepreis je Dose", bold=True)
put(ws, 14, 2, "=B12", fmt=EUR2, bold=True)
put(ws, 15, 1, "Ø Rohertrag je Dose", bold=True)
c = put(ws, 15, 2, "=D12", fmt=EUR2, bold=True, size=12)
c.fill = YELLOW; c.border = BOX

note(ws, 8, 9, "Preise netto — ohne 20 % USt und ohne die 25 Cent Pfand (durchlaufender Posten).")
note(ws, 10, 9, "Faustregel: Gastro und Direktverkauf verdienen das Geld, der LEH macht den Umsatz.")
ws.row_dimensions[10].height = 28

# ============================================================== BREAK-EVEN
ws = wb.create_sheet("Break-even")
title(ws, "BREAK-EVEN DER ERSTEN CHARGE", 8)
widths(ws, {"A": 42, "B": 19, "C": 19, "D": 22, "E": 19, "F": 20, "G": 19, "H": 3, "I": 48})

section(ws, 3, "ZU DECKENDE KOSTEN", 3)
put(ws, 4, 1, "Einmalkosten gesamt")
put(ws, 4, 2, "=Annahmen!$B$33", fmt=EUR2, color=GREEN)
put(ws, 5, 1, "Fixkosten Jahr 1")
put(ws, 5, 2, "=Annahmen!$B$40", fmt=EUR2, color=GREEN)
put(ws, 6, 1, "Summe", bold=True)
put(ws, 6, 2, "=B4+B5", fmt=EUR2, bold=True)
ws["A6"].border = TOPLINE; ws["B6"].border = TOPLINE

section(ws, 8, "BREAK-EVEN", 3)
put(ws, 9, 1, "Ø Rohertrag je Dose")
put(ws, 9, 2, "='Marge & Preiskette'!$B$15", fmt=EUR2, color=GREEN)
put(ws, 10, 1, "BREAK-EVEN (verkaufte Dosen)", bold=True)
c = put(ws, 10, 2, "=IFERROR($B$6/$B$9,0)", fmt=NUM, bold=True, size=12)
c.fill = YELLOW; c.border = BOX
put(ws, 11, 1, "Chargengröße")
put(ws, 11, 2, "=Annahmen!$B$4", fmt=NUM, color=GREEN)
put(ws, 12, 1, "Break-even in % der Charge", bold=True)
put(ws, 12, 2, "=IFERROR($B$10/$B$11,0)", fmt=PCT, bold=True)
put(ws, 12, 9, '=IF(B12>1,"ACHTUNG: Die erste Charge kann die Kosten rechnerisch nicht decken.","OK — der Break-even liegt innerhalb der ersten Charge.")',
    bold=True, color="C00000", wrap=True)
ws.row_dimensions[12].height = 30

section(ws, 14, "SZENARIEN CHARGE 1 — die Ware wird voll produziert, aber nur teilweise verkauft", 7)
header_row(ws, 15, ["Abverkauf", "Verkaufte Dosen", "Umsatz netto",
                    "Wareneinsatz (volle Charge)", "Deckungsbeitrag",
                    "Einmal- + Fixkosten", "ERGEBNIS"])
for i, share in enumerate([0.50, 0.75, 1.00]):
    rr = 16 + i
    put(ws, rr, 1, share, fmt=PCT, color=BLUE, inp=True)
    put(ws, rr, 2, f"=$B$11*A{rr}", fmt=NUM)
    put(ws, rr, 3, f"=B{rr}*'Marge & Preiskette'!$B$14", fmt=EUR2)
    put(ws, rr, 4, f"=-$B$11*'Marge & Preiskette'!$B$5", fmt=EUR2)
    put(ws, rr, 5, f"=C{rr}+D{rr}", fmt=EUR2)
    put(ws, rr, 6, "=-$B$6", fmt=EUR2)
    c = put(ws, rr, 7, f"=E{rr}+F{rr}", fmt=EUR2, bold=True)
    c.fill = MID
note(ws, 16, 9, "Nicht verkaufte Dosen kosten trotzdem Geld — deshalb wird der Wareneinsatz "
                "immer für die volle Charge gerechnet. Das ist der ehrliche Blick.")
ws.row_dimensions[16].height = 44

# ========================================================= JAHRESSZENARIEN
ws = wb.create_sheet("Jahresszenarien")
title(ws, "JAHRESSZENARIEN — ab wann trägt es dich?", 9)
widths(ws, {"A": 24, "B": 18, "C": 17, "D": 16, "E": 18, "F": 18, "G": 18, "H": 16,
            "I": 3, "J": 46})

header_row(ws, 3, ["Szenario", "Absatz Dosen / Jahr", "Vollkosten € / Dose", "Ø Abgabepreis",
                   "Rohertrag gesamt", "Fixkosten p. a.", "ERGEBNIS", "je Monat"])
for i, (name, vol, fixk) in enumerate([("Nebenbei", 20000, 4000), ("Halbtags", 80000, 25000),
                                       ("Hauptberuf", 200000, 70000), ("Ernsthaft", 500000, 180000)]):
    rr = 4 + i
    put(ws, rr, 1, name, bold=True)
    put(ws, rr, 2, vol, fmt=NUM, color=BLUE, inp=True)
    put(ws, rr, 3, "=IFERROR(INDEX('Stückkosten'!$B$26:$E$28,"
                   f"MATCH(B{rr},'Stückkosten'!$A$26:$A$28,1),"
                   "MATCH(Annahmen!$B$6,'Stückkosten'!$B$25:$E$25,0)),0)", fmt=EUR3)
    put(ws, rr, 4, "='Marge & Preiskette'!$B$14", fmt=EUR2, color=GREEN)
    put(ws, rr, 5, f"=B{rr}*(D{rr}-C{rr})", fmt=EUR2)
    put(ws, rr, 6, fixk, fmt=EUR2, color=BLUE, inp=True)
    c = put(ws, rr, 7, f"=E{rr}-F{rr}", fmt=EUR2, bold=True)
    c.fill = MID
    put(ws, rr, 8, f"=G{rr}/12", fmt=EUR2)

note(ws, 4, 10, "Fixkosten sind Eingaben — schätze sie ehrlich. Ab 'Hauptberuf' gehört dein eigenes Gehalt hinein, sonst rechnest du dich reich.")
ws.row_dimensions[4].height = 44
note(ws, 6, 10, "Ab 100.000 Dosen bleibt die Kostenstufe rechnerisch stehen. Bei 500.000 wären die echten Stückkosten niedriger — hier bewusst konservativ.")
ws.row_dimensions[6].height = 44

section(ws, 10, "WAS DAS HAUPTBERUF-SZENARIO PRAKTISCH BEDEUTET", 8)
put(ws, 11, 1, "Dosen je Tag")
put(ws, 11, 2, "=B6/365", fmt=NUM)
put(ws, 12, 1, "Gastro-Betriebe nötig")
put(ws, 12, 2, 100, fmt=NUM, color=BLUE, inp=True)
put(ws, 13, 1, "Dosen je Betrieb und Tag")
put(ws, 13, 2, "=IFERROR($B$11/$B$12,0)", fmt=NUM1)
note(ws, 13, 10, "Das ist die eigentliche Arbeit: so viele Betriebe gewinnen UND halten.")

# ============================================================== VERKOSTUNG
ws = wb.create_sheet("Verkostung")
title(ws, "BLINDVERKOSTUNG — AUSWERTUNG", 11)
widths(ws, {"A": 34, "B": 12, "C": 12, "D": 12, "E": 12, "F": 12, "G": 12, "H": 12,
            "I": 12, "J": 12, "K": 38, "L": 46})

section(ws, 3, "SO GEHT ES", 11)
for i, t in enumerate([
    "1. Muster in identische, undurchsichtige Becher füllen und NUR mit dem Zahlencode beschriften.",
    "2. Reihenfolge bei jedem Teilnehmer wechseln. Wasser und etwas Brot zum Neutralisieren dazustellen.",
    "3. Erst den Papierbogen ausfüllen lassen (PDF im Projektordner), danach die Zahlen hier eintragen.",
    "4. Die Legende unten darf niemand sehen, bevor alle Bögen ausgefüllt sind.",
]):
    put(ws, 4 + i, 1, t, size=9)
    ws.merge_cells(start_row=4 + i, start_column=1, end_row=4 + i, end_column=8)

section(ws, 9, "BLINDCODES — Legende (geheim halten!)", 6)
header_row(ws, 10, ["Code", "Muster"])
for i, (code, name) in enumerate([(471, "Referenz — Fresca Original (Import)"),
                                  (628, "Variante A — Vollzucker"),
                                  (359, "Variante B — Fresca-Original (Hybrid, mit Süßstoff)"),
                                  (815, "Variante C — Nur Zucker, reduziert")]):
    put(ws, 11 + i, 1, code, fmt=NUM, color=BLUE, inp=True, align="center")
    put(ws, 11 + i, 2, name, color=BLUE)
    ws.merge_cells(start_row=11 + i, start_column=2, end_row=11 + i, end_column=6)

section(ws, 16, "AUSWERTUNG — rechnet automatisch", 6)
header_row(ws, 17, ["Kennzahl", "Code 1", "Code 2", "Code 3", "Code 4"])
for i in range(4):
    put(ws, 17, 2 + i, f"=A{11+i}", bold=True, align="center", color=GREEN)
for rr, lab in [(18, "Ø Gesamteindruck (1–10)"),
                (19, "Ø Süßeempfinden (1 = zu wenig … 9 = zu süß)"),
                (20, "Anzahl \"würde ich kaufen\""),
                (21, "Anteil Kaufpräferenz"),
                (22, "Differenz zur Referenz (Gesamteindruck)")]:
    put(ws, rr, 1, lab)
for i in range(4):
    gcol = get_column_letter(2 + i)   # B..E  Gesamteindruck-Daten
    scol = get_column_letter(6 + i)   # F..I  Süße-Daten
    put(ws, 18, 2 + i, f"=IFERROR(AVERAGE({gcol}$30:{gcol}$77),0)", fmt=NUM1, bold=True)
    put(ws, 19, 2 + i, f"=IFERROR(AVERAGE({scol}$30:{scol}$77),0)", fmt=NUM1)
    put(ws, 20, 2 + i, f"=COUNTIF($J$30:$J$77,{gcol}$17)", fmt=NUM)
    put(ws, 21, 2 + i, f"=IFERROR({gcol}20/$B$24,0)", fmt=PCT)
    put(ws, 22, 2 + i, f"={gcol}18-$B$18", fmt=NUM1)

put(ws, 24, 1, "Teilnehmer erfasst", bold=True)
put(ws, 24, 2, "=COUNT($B$30:$B$77)", fmt=NUM, bold=True)
put(ws, 25, 1, "Bestes Muster (Ø Gesamteindruck)", bold=True)
c = put(ws, 25, 2, "=IF($B$24=0,\"— noch keine Daten erfasst —\","
                   "IFERROR(INDEX($B$11:$B$14,MATCH(INDEX($B$17:$E$17,"
                   "MATCH(MAX($B$18:$E$18),$B$18:$E$18,0)),$A$11:$A$14,0)),\"—\"))", bold=True)
c.fill = YELLOW; c.border = BOX
ws.merge_cells(start_row=25, start_column=2, end_row=25, end_column=7)
put(ws, 26, 1, "ENTSCHEIDUNGSREGEL", bold=True, color="C00000")
put(ws, 26, 2, "Schlägt keine eigene Variante die Referenz um mindestens 0,5 Punkte, ist die Rezeptur noch nicht fertig.",
    size=9, italic=True, color="C00000")
ws.merge_cells(start_row=26, start_column=2, end_row=26, end_column=11)

section(ws, 27, "DATENEINGABE — eine Zeile je Teilnehmer", 11)
header_row(ws, 28, ["Teilnehmer", "Gesamt Code 1", "Gesamt Code 2", "Gesamt Code 3",
                    "Gesamt Code 4", "Süße Code 1", "Süße Code 2", "Süße Code 3",
                    "Süße Code 4", "Kauf-Code", "Notiz"])
for i, v in enumerate(["BEISPIEL", 7, 5, 8, 6, 5, 8, 5, 3, 359, "B am rundesten, C zu trocken"]):
    c = put(ws, 29, 1 + i, v, italic=True, color=GREY_T, size=9)
    c.fill = GREYF
put(ws, 29, 12, "◄ Beispielzeile. Sie wird NICHT mitgerechnet — echte Daten ab Zeile 30.",
    italic=True, size=9, color="C00000")

for i in range(1, 49):
    rr = 29 + i
    put(ws, rr, 1, i, fmt=NUM, align="center")
    for col in range(2, 12):
        cc = ws.cell(row=rr, column=col)
        cc.font = Font(name=FONT, size=10, color=BLUE)
        cc.fill = YELLOW
        cc.border = BOX
        if col <= 10:
            cc.number_format = NUM

for w in wb.worksheets:
    w.sheet_view.showGridLines = False
    w.freeze_panes = "A2"

wb.save(OUT)
print("gespeichert:", OUT)
